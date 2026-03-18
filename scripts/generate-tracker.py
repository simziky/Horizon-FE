"""
Generates the Horizon FE Optimization Tracker Excel file.
Run: python scripts/generate-tracker.py
"""

from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()

# ─────────────────────────────────────────────
# COLOUR PALETTE
# ─────────────────────────────────────────────
C_HEADER_DARK   = "1E293B"   # slate-800  – main header bg
C_HEADER_MID    = "334155"   # slate-700  – phase header bg
C_PHASE1_ACCENT = "0F172A"   # deep navy  – phase 1 stripe
C_PHASE2_ACCENT = "1E3A5F"   # deep blue  – phase 2 stripe
C_DONE          = "DCFCE7"   # green-100
C_DONE_FONT     = "166534"   # green-800
C_PARTIAL       = "FEF9C3"   # yellow-100
C_PARTIAL_FONT  = "854D0E"   # yellow-800
C_TODO          = "FEE2E2"   # red-100
C_TODO_FONT     = "991B1B"   # red-800
C_PLANNED       = "EDE9FE"   # violet-100
C_PLANNED_FONT  = "5B21B6"   # violet-800
C_ROW_ALT       = "F8FAFC"   # slate-50
C_WHITE         = "FFFFFF"
C_GRID          = "CBD5E1"   # slate-300

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border(top=True, bottom=True, left=True, right=True):
    s = Side(style="thin", color=C_GRID)
    n = Side(style=None)
    return Border(
        top=s if top else n,
        bottom=s if bottom else n,
        left=s if left else n,
        right=s if right else n,
    )

def bold_border():
    s = Side(style="medium", color="94A3B8")
    return Border(top=s, bottom=s, left=s, right=s)

# ─────────────────────────────────────────────
# SHEET 1 — TRACKER
# ─────────────────────────────────────────────
ws = wb.active
ws.title = "Tracker"

# ── Column widths ──────────────────────────────
col_widths = {
    "A": 5,   # #
    "B": 38,  # Task / Sub-task
    "C": 18,  # Area
    "D": 16,  # Status
    "E": 14,  # Priority
    "F": 14,  # Effort
    "G": 28,  # Key File(s)
    "H": 30,  # Notes / Next Action
    "I": 14,  # Owner
    "J": 14,  # Target Week
    "K": 14,  # Done Date
}
for col, width in col_widths.items():
    ws.column_dimensions[col].width = width

ws.row_dimensions[1].height = 36
ws.row_dimensions[2].height = 22

# ── Title row ─────────────────────────────────
ws.merge_cells("A1:K1")
title_cell = ws["A1"]
title_cell.value = "Horizon FE — Optimization Tracker  |  Phase 1 & 2"
title_cell.font = Font(name="Calibri", bold=True, size=16, color=C_WHITE)
title_cell.fill = fill(C_HEADER_DARK)
title_cell.alignment = Alignment(horizontal="center", vertical="center")

# ── Column headers ────────────────────────────
headers = ["#", "Task / Sub-task", "Area", "Status", "Priority",
           "Effort", "Key File(s)", "Notes / Next Action", "Owner",
           "Target Week", "Done Date"]
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=2, column=col_idx, value=h)
    cell.font = Font(name="Calibri", bold=True, size=10, color=C_WHITE)
    cell.fill = fill(C_HEADER_MID)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border()

# ── Freeze panes ──────────────────────────────
ws.freeze_panes = "C3"

# ── Status badge helper ───────────────────────
STATUS_STYLES = {
    "✅ Done":      (C_DONE,    C_DONE_FONT),
    "⚠️ Partial":  (C_PARTIAL, C_PARTIAL_FONT),
    "❌ Not Done":  (C_TODO,    C_TODO_FONT),
    "📋 Planned":  (C_PLANNED, C_PLANNED_FONT),
    "🚫 Ignored":  ("E2E8F0",  "64748B"),
}

def write_row(ws, row_num, data, is_phase_header=False, is_section=False, alt=False):
    """
    data = [num, task, area, status, priority, effort, files, notes, owner, week, done]
    """
    bg = C_WHITE
    if is_phase_header:
        bg = C_PHASE1_ACCENT if "Phase 1" in str(data[1]) else C_PHASE2_ACCENT
    elif is_section:
        bg = "E2E8F0"
    elif alt:
        bg = C_ROW_ALT

    for col_idx, val in enumerate(data, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=val)
        cell.border = thin_border()
        cell.alignment = Alignment(vertical="center", wrap_text=True)

        if is_phase_header:
            cell.font = Font(name="Calibri", bold=True, size=11, color=C_WHITE)
            cell.fill = fill(bg)
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif is_section:
            cell.font = Font(name="Calibri", bold=True, size=10, color="1E293B")
            cell.fill = fill(bg)
        else:
            cell.font = Font(name="Calibri", size=9, color="1E293B")
            cell.fill = fill(bg)

        # Status colouring
        if col_idx == 4 and val in STATUS_STYLES:
            bg_s, fg_s = STATUS_STYLES[val]
            cell.fill = fill(bg_s)
            cell.font = Font(name="Calibri", bold=True, size=9, color=fg_s)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Priority colouring
        if col_idx == 5:
            p_colors = {"🔴 Critical": ("FECACA","991B1B"),
                        "🟠 High":     ("FED7AA","9A3412"),
                        "🟡 Medium":   ("FEF9C3","854D0E"),
                        "🟢 Low":      ("DCFCE7","166534")}
            if val in p_colors:
                bg_p, fg_p = p_colors[val]
                cell.fill = fill(bg_p)
                cell.font = Font(name="Calibri", bold=True, size=9, color=fg_p)
                cell.alignment = Alignment(horizontal="center", vertical="center")

        # Number / effort / week centred
        if col_idx in (1, 6, 10, 11):
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[row_num].height = 40 if is_phase_header else (18 if is_section else 36)

# ─────────────────────────────────────────────
# DATA
# ─────────────────────────────────────────────
rows = []

# ── PHASE 1 HEADER ────────────────────────────
rows.append(("", "⚡  PHASE 1 — Critical Performance Issues  (Weeks 1–4)", "", "", "", "", "", "", "", "", "", "PHASE"))

# ═══ SECTION 1.0 ─ Folder Restructure (Bonus) ═
rows.append(("", "📁  Folder Restructure & Code Organisation  [BONUS — User + AI]", "", "", "", "", "", "", "", "", "", "SECTION"))
rows.append(("1.0", "Design target folder structure & migration map",
             "Architecture", "✅ Done", "🔴 Critical", "M",
             "docs/folder-structure-refactor.md",
             "Full target tree documented with migration map",
             "", "Week 1", "Feb 2026", "ROW"))
rows.append(("1.1", "Move utils/ UI components → components/ui/ (Loader, CircularLoader, ProgressLoader, AntdComponents)",
             "Architecture", "✅ Done", "🔴 Critical", "S",
             "components/ui/Loader.tsx\ncomponents/ui/CircularLoader.tsx\ncomponents/ui/ProgressLoader.tsx\ncomponents/ui/AntdComponents.tsx",
             "Files moved and imports updated",
             "", "Week 1", "Feb 2026", "ROW"))
rows.append(("1.2", "Move lib/countries.ts → constants/countries.ts",
             "Architecture", "✅ Done", "🟡 Medium", "XS",
             "constants/countries.ts",
             "Static data separated from lib/",
             "", "Week 1", "Feb 2026", "ROW"))
rows.append(("1.3", "Move all app/(dashboard)/_components/ → components/ui/ + components/layout/ + components/modals/",
             "Architecture", "✅ Done", "🔴 Critical", "L",
             "components/layout/DashNav.tsx\ncomponents/layout/DashSider.tsx\ncomponents/modals/RenewSubModal.tsx",
             "All shared dashboard UI components relocated",
             "", "Week 1", "Feb 2026", "ROW"))
rows.append(("1.4", "Move all feature _components/ → components/features/<feature>/",
             "Architecture", "✅ Done", "🔴 Critical", "XL",
             "components/features/dashboard/\ncomponents/features/keepa/\ncomponents/features/seller/\ncomponents/features/upc-scanner/",
             "Feature-scoped components now live in components/features/",
             "", "Week 1", "Feb 2026", "ROW"))
rows.append(("1.5", "Create new components/modals/ (AmazonConnectModal, ExpiredSubscriptionModal, PackageRestrictionModal)",
             "Architecture", "✅ Done", "🟠 High", "S",
             "components/modals/AmazonConnectModal.tsx\ncomponents/modals/ExpiredSubscriptionModal.tsx\ncomponents/modals/PackageRestrictionModal.tsx",
             "Shared modals centralised",
             "", "Week 1", "Feb 2026", "ROW"))
rows.append(("1.6", "Rename redux/ → store/ and redux/slice/ → store/slices/",
             "Architecture", "🚫 Ignored", "🟢 Low", "L",
             "redux/ (all files)",
             "Assessed and deliberately excluded. High churn with no user-facing benefit; all @/redux/* imports work correctly as-is.",
             "", "Week 4", "", "ROW"))
rows.append(("1.7", "Remove duplicate (auth1)/ route group; merge into (auth)/",
             "Architecture", "✅ Done", "🟡 Medium", "M",
             "app/(auth1)/\napp/(auth)/",
             "All 6 auth1 files deleted (Login, Steps, layout, auth/page, signup/Signup, signup/page). No remaining imports or references to auth1 anywhere in the codebase. Only (auth)/ route group active.",
             "", "Week 3", "Mar 2026", "ROW"))

# ═══ SECTION 1 ─ Virtual Scrolling ═══════════
rows.append(("", "📜  1. Virtual Scrolling for Large Lists", "", "", "", "", "", "", "", "", "", "SECTION"))
rows.append(("1.8", "Research & library decision (react-window vs react-virtualized)",
             "Virtual Scroll", "✅ Done", "🔴 Critical", "S",
             "docs/virtual-scrolling-technical-spec.md",
             "react-window selected; UPC results = highest priority target",
             "", "Week 1", "Feb 2026", "ROW"))
rows.append(("1.9", "UPC Scanner results table + details table: virtual scroll + perf optimizations",
             "Virtual Scroll", "✅ Done", "🔴 Critical", "L",
             "components/features/upc-scanner/scan-results-table.tsx\ncomponents/features/upc-scanner/scan-details-table.tsx\napp/globals.css",
             "scan-results-table: Ant Design virtual prop, scroll jank fix, memoized tableData/columns/handlers, memory fix. scan-details-table: fully rewritten from dual custom HTML tables to Ant Design Table with virtual prop, fixed:left sticky columns (UPC/EAN, Product Cost, Buy Box Price), Ant Design native sorters replace manual sort state.",
             "", "Week 2", "Mar 2026", "ROW"))
rows.append(("1.10", "Library decision: react-window vs Ant Design native virtual",
             "Virtual Scroll", "✅ Done", "🔴 Critical", "XS",
             "components/features/upc-scanner/scan-results-table.tsx",
             "Ant Design v5 native `virtual` prop used — react-window not required. Avoids extra dependency; scroll jank and memory issues resolved separately.",
             "", "Week 2", "Mar 2026", "ROW"))
rows.append(("1.11", "Benchmark virtual scroll: 100 / 500 / 1k / 5k / 10k mock rows",
             "Virtual Scroll", "✅ Done", "🟠 High", "M",
             "app/upc-benchmark/page.tsx\napp/upc-benchmark/UpcBenchmarkClient.tsx",
             "Benchmark page live at /upc-benchmark. Measures mount time, DOM row count (virtual proof), filter P95, heap used. React Profiler + window.__upcBenchResult for automation.",
             "", "Week 2", "Mar 2026", "ROW"))
rows.append(("1.12", "Validate interaction parity (details, restart, delete, download actions in virtual rows)",
             "Virtual Scroll", "✅ Done", "🔴 Critical", "M",
             "components/features/upc-scanner/scan-results-table.tsx",
             "All actions (view details, restart scan, delete, download Excel) verified working inside virtual rows. useMemo on columns preserves stable action handler refs.",
             "", "Week 2", "Mar 2026", "ROW"))
rows.append(("1.13", "Seller products list: add react-window VariableSizeList (gate: perPage ≥ 25 or infinite scroll)",
             "Virtual Scroll", "🚫 Ignored", "🟢 Low", "L",
             "components/features/seller/Seller.tsx",
             "Assessed and deliberately excluded. perPage=10 — virtualisation threshold not met. LazyKeepaChart already handles the main render cost per card.",
             "", "Week 4+", "", "ROW"))

# ═══ SECTION 2 ─ API Caching ═══════════════
rows.append(("", "⚡  2. API Caching & Request Optimization", "", "", "", "", "", "", "", "", "", "SECTION"))
rows.append(("2.1", "Design caching strategy & TTL spec",
             "API Caching", "✅ Done", "🔴 Critical", "S",
             "docs/caching-strategy-technical-spec.md",
             "Full TTL matrix, cache keys, invalidation rules documented",
             "", "Week 1", "Feb 2026", "ROW"))
rows.append(("2.2", "Keepa endpoints: set keepUnusedDataFor (hard TTL) per endpoint",
             "API Caching", "✅ Done", "🔴 Critical", "S",
             "redux/api/keepa.ts",
             "keepa_summary=1200s, price_history/sales_rank/ratings=1800s",
             "", "Week 2", "Feb 2026", "ROW"))
rows.append(("2.3", "IP Alert endpoint: set keepUnusedDataFor=600 (hard TTL 10 min)",
             "API Caching", "✅ Done", "🔴 Critical", "S",
             "redux/api/productsApi.ts",
             "keepUnusedDataFor: 600 set at line 125",
             "", "Week 2", "Feb 2026", "ROW"))
rows.append(("2.4", "Keepa endpoints: set refetchOnMountOrArgChange (soft TTL) at API level",
             "API Caching", "✅ Done", "🔴 Critical", "S",
             "redux/api/keepa.ts",
             "refetchOnMountOrArgChange: 900 (15 min) added at createApi level. Hard TTLs (keepUnusedDataFor) remain per-endpoint: 1200s summary, 1800s others. refetchOnReconnect: true added.",
             "", "Week 2", "Mar 2026", "ROW"))
rows.append(("2.5", "Fix remaining APIs still using refetchOnMountOrArgChange: 10 (too aggressive)",
             "API Caching", "✅ Done", "🟠 High", "M",
             "redux/api/auth.ts\nredux/api/monitorApi.ts\nredux/api/sellerApi.ts\nredux/api/user.ts",
             "All four APIs raised from 10s → 300s (5 min). Tag invalidation handles mutation-driven freshness. All endpoints in these APIs are lazy queries so automatic refetch only triggers on reconnect.",
             "", "Week 2", "Mar 2026", "ROW"))
rows.append(("2.6", "Add cache key normalisation helpers (ASIN uppercase, period normalised to 7d/30d/90d/1y/all)",
             "API Caching", "✅ Done", "🟠 High", "M",
             "utils/cacheKeys.ts\nredux/api/keepa.ts",
             "Created utils/cacheKeys.ts: normalizeAsin(), normalizePeriod(), normalizeMarketplaceId(), keepaCacheKey(). All 4 keepa serializeQueryArgs updated to use keepaCacheKey(). ratingReview period normalisation bug fixed (was using period ?? '' without normalization).",
             "", "Week 2", "Mar 2026", "ROW"))
rows.append(("2.7", "Optional L2 localStorage cache for Keepa heavy reads (cold-start speed)",
             "API Caching", "🚫 Ignored", "🟡 Medium", "L",
             "utils/ (new file: keepaLocalCache.ts)",
             "Ignored: RTK Query hard TTL (20-30 min) already covers in-session caching. Cold-start gain is minimal; adds localStorage stale-data risk and storage management overhead.",
             "", "—", "", "ROW"))
rows.append(("2.8", "Add dev-mode cache observability logging (hit / miss / stale)",
             "API Caching", "🚫 Ignored", "🟢 Low", "S",
             "redux/queryInterceptor.tsx",
             "Ignored: Redux DevTools extension already exposes full RTK Query cache state, query status and timing. Duplicate tooling.",
             "", "—", "", "ROW"))
rows.append(("2.9", "Request deduplication (RTK Query native — verify no double triggers)",
             "API Caching", "✅ Done", "🟠 High", "XS",
             "redux/api/ (all)",
             "RTK Query deduplicates by default. Audit lazy query triggers for duplicate fires",
             "", "Week 1", "Feb 2026", "ROW"))
rows.append(("2.10", "Request batching for multi-product queries",
             "API Caching", "🚫 Ignored", "🟡 Medium", "XL",
             "redux/api/productsApi.ts",
             "Ignored: requires BE support or client-side queue. UI is single-product per request by design. No user-visible gain without BE change.",
             "", "—", "", "ROW"))
rows.append(("2.11", "Optimistic UI updates for user actions",
             "API Caching", "🚫 Ignored", "🟡 Medium", "L",
             "redux/api/",
             "Ignored: monitor/unmonitor and settings mutations are infrequent and low-latency. Tag invalidation → refetch is clean and correct. Optimistic updates add rollback complexity for negligible UX gain.",
             "", "—", "", "ROW"))

# ═══ SECTION 3 ─ Chart Rendering ══════════════
rows.append(("", "📊  3. Chart Rendering Optimization", "", "", "", "", "", "", "", "", "", "SECTION"))
rows.append(("3.1", "Chart benchmark harness setup (recharts baseline at 300/1k/5k points)",
             "Charts", "✅ Done", "🔴 Critical", "M",
             "scripts/benchmark-charts.js\napp/chart-benchmark/\ndocs/chart-benchmark-results.json",
             "Results: 300pt=72.9ms mount, 5k=124.9ms mount, 5k update p95=82.5ms",
             "", "Week 1", "Feb 2026", "ROW"))
rows.append(("3.2", "Chart library evaluation (recharts SVG vs Chart.js canvas vs custom SVG)",
             "Charts", "✅ Done", "🔴 Critical", "M",
             "docs/chart-library-evaluation.md\ncomponents/features/keepa/KeepaLineChart.tsx",
             "Revised decision: custom zero-dependency SVG chart built (KeepaLineChart.tsx) — lighter than Chart.js, no external dependency, full rendering control. recharts removed from KeepaChart entirely.",
             "", "Week 1", "Mar 2026", "ROW"))
rows.append(("3.3", "Build custom SVG chart component for KeepaChart (replaces recharts + Chart.js POC)",
             "Charts", "✅ Done", "🔴 Critical", "XL",
             "components/features/keepa/KeepaLineChart.tsx\ncomponents/features/keepa/KeepaChart.tsx",
             "Zero-dep KeepaLineChart.tsx built: dual Y axes, flatline on null data, RAF-throttled crosshair sync, niceRange() ticks, domain padding, per-line pixelOffset, useLayoutEffect width, top/bottom edge label guards, tooltip render prop. recharts bundle (~400KB) eliminated.",
             "", "Week 2-3", "Mar 2026", "ROW"))
rows.append(("3.4", "Benchmark recharts vs Chart.js side-by-side comparison",
             "Charts", "🚫 Ignored", "🟡 Medium", "M",
             "scripts/benchmark-charts.js",
             "Superseded. Custom SVG KeepaLineChart implemented directly — Chart.js POC no longer needed. Decision already made and validated.",
             "", "Week 4+", "", "ROW"))
rows.append(("3.5", "Migrate KeepaChart.tsx off recharts to custom SVG",
             "Charts", "✅ Done", "🟠 High", "XL",
             "components/features/keepa/KeepaChart.tsx",
             "Full migration complete: Map pre-indexing O(1) lookup, Intl.DateTimeFormat memoized per range, invalid date guards, memoized line/axis/data configs, synchronized crosshair via syncedTimestamp, abbreviateNumber handles negatives, calcDomain clamps ≥0 data to avoid negative ticks.",
             "", "Week 3-4", "Mar 2026", "ROW"))
rows.append(("3.6", "Migrate seller/keepa-chart.tsx to custom SVG (conditional on card density growth)",
             "Charts", "🚫 Ignored", "🟡 Medium", "L",
             "components/features/seller/keepa-chart.tsx",
             "Assessed and deliberately excluded. LazyKeepaChart IntersectionObserver gate already resolves the render cost. perPage=10 — recharts overhead acceptable at this density.",
             "", "Week 4+", "", "ROW"))
rows.append(("3.7", "Web Workers for heavy Keepa price calculations & date formatting",
             "Charts", "🚫 Ignored", "🟠 High", "XL",
             "components/features/keepa/KeepaChart.tsx",
             "Assessed and deliberately excluded. Map pre-indexing (O(n²)→O(1)) + useMemo already resolves main thread pressure. Web Worker adds large complexity for diminishing gain at current dataset sizes.",
             "", "Week 3", "", "ROW"))
rows.append(("3.8", "RAF-throttled hover sync across charts (prevent UI jank)",
             "Charts", "✅ Done", "🟠 High", "M",
             "components/features/keepa/KeepaLineChart.tsx",
             "requestAnimationFrame throttle on mousemove inside KeepaLineChart. Single syncedTimestamp string state drives crosshair sync across all 3 charts. No additional debounce needed for hover.",
             "", "Week 2", "Mar 2026", "ROW"))

# ═══ SECTION 4 ─ Code Splitting ═══════════════
rows.append(("", "🧩  4. Code Splitting & Bundle Management", "", "", "", "", "", "", "", "", "", "SECTION"))
rows.append(("4.1", "Install & configure @next/bundle-analyzer",
             "Code Splitting", "✅ Done", "🔴 Critical", "S",
             "next.config.ts\npackage.json",
             "ANALYZE=true flag enables the analyzer; configured in next.config.ts",
             "", "Week 1", "Feb 2026", "ROW"))
rows.append(("4.2", "Run bundle analysis — identify chunks > 50KB",
             "Code Splitting", "✅ Done", "🔴 Critical", "S",
             "docs/ (analysis output)",
             "Analysis run as part of Week 1 plan",
             "", "Week 1", "Feb 2026", "ROW"))
rows.append(("4.3", "ProductDetails: dynamic() imports for all sub-components (BuyboxAnalysis, MarketAnalysis, ProfitabilityCalculator, etc.)",
             "Code Splitting", "✅ Done", "🔴 Critical", "L",
             "components/features/dashboard/ProductDetails.tsx",
             "Nav, ProductOverview, BuyboxAnalysis, MarketAnalysis, ProfitabilityCalculator, ProductEligibility, TopSellers, CalculationResults all lazy-loaded",
             "", "Week 2", "Feb 2026", "ROW"))
rows.append(("4.4", "KeepaChart: dynamic() import with SSR disabled",
             "Code Splitting", "✅ Done", "🔴 Critical", "S",
             "app/(dashboard)/keepa/page.tsx",
             "ssr: false + dynamic import applied",
             "", "Week 2", "Feb 2026", "ROW"))
rows.append(("4.5", "Go-Compare Navigation: dynamic() import",
             "Code Splitting", "✅ Done", "🟠 High", "S",
             "app/(dashboard)/go-compare/layout.tsx",
             "Navigation lazy-loaded in layout",
             "", "Week 2", "Feb 2026", "ROW"))
rows.append(("4.6", "Dashboard layout modals: dynamic() imports (ExpiredSub, AmazonConnect, PackageRestriction)",
             "Code Splitting", "✅ Done", "🟠 High", "S",
             "app/(dashboard)/layout.tsx",
             "Three modal components lazy-loaded in layout",
             "", "Week 2", "Feb 2026", "ROW"))
rows.append(("4.7", "Seller page: lazy-load heavy components (KeepaChart in cards)",
             "Code Splitting", "✅ Done", "🟠 High", "M",
             "components/features/seller/LazyKeepaChart.tsx\ncomponents/features/seller/Seller.tsx",
             "Created LazyKeepaChart.tsx wrapper: IntersectionObserver (threshold 0.1) gates chart mount until card enters viewport. Once in view, never unmounts. Shows pulse skeleton until visible. Removed redundant dynamic() in Seller.tsx.",
             "", "Week 3", "Mar 2026", "ROW"))
rows.append(("4.8", "UPC Scanner page: lazy-load scan result sub-sections",
             "Code Splitting", "🚫 Ignored", "🟡 Medium", "M",
             "components/features/upc-scanner/UpcScanner.tsx",
             "Assessed and deliberately excluded. scan-details-table already virtualised; further lazy chunking adds complexity with minimal gain.",
             "", "Week 3", "", "ROW"))
rows.append(("4.9", "Product page memoization: React.memo + useMemo for ProductPage wrapper",
             "Code Splitting", "✅ Done", "🟠 High", "S",
             "app/(dashboard)/dashboard/product/[asin]/page.tsx",
             "React.memo and useMemo for productDetailsProps applied",
             "", "Week 2", "Feb 2026", "ROW"))

# ── PHASE 2 HEADER ────────────────────────────
rows.append(("", "🎨  PHASE 2 — User Experience Improvements  (Weeks 5–8)", "", "", "", "", "", "", "", "", "", "PHASE"))

# ═══ SECTION 5 ─ Progressive Loading ══════════
rows.append(("", "⏳  5. Progressive Loading & Suspense", "", "", "", "", "", "", "", "", "", "SECTION"))
rows.append(("5.1", "Add React.Suspense boundaries at route/layout level",
             "Loading UX", "✅ Done", "🟠 High", "M",
             "app/(dashboard)/layout.tsx\napp/(dashboard)/keepa/page.tsx\napp/(dashboard)/go-compare/layout.tsx",
             "Dashboard layout has <Suspense fallback={pulse}>{children}</Suspense> at line 113 covering all route transitions. KeepaChart, Navigation (GoCompare), MarketAnalysis, ProfitabilityCalculator all have loading: fallbacks on dynamic() — equivalent to Suspense boundary. No gaps found.",
             "", "Week 5", "Mar 2026", "ROW"))
rows.append(("5.2", "Progressive loading for Seller details: show basic card info first, load KeepaChart deferred",
             "Loading UX", "✅ Done", "🟠 High", "L",
             "components/features/seller/Seller.tsx",
             "Implemented via LazyKeepaChart.tsx: IntersectionObserver (threshold 0.1) gates KeepaChart mount until card enters viewport. Pulse skeleton shown until in view. Never re-mounts once loaded.",
             "", "Week 5", "Mar 2026", "ROW"))
rows.append(("5.3", "Progressive loading for ProductDetails: show QuickInfo immediately, defer analytics panels",
             "Loading UX", "🚫 Ignored", "🟠 High", "M",
             "components/features/dashboard/ProductDetails.tsx",
             "Assessed and deliberately excluded. FinalLoader is a conscious full-page loader pattern — all 3 APIs (getItem, getBuyboxDetails, getIpAlert) resolve before any content shows. MarketAnalysis/ProfitabilityCalculator have dynamic() skeletons but they never fire since data is ready. Partial render would create inconsistent UI.",
             "", "Week 5", "", "ROW"))
rows.append(("5.4", "DashNav notifications: infinite scroll with IntersectionObserver + paginated API",
             "Loading UX", "✅ Done", "🟠 High", "M",
             "components/layout/DashNav.tsx\nredux/api/user.ts",
             "getNotifications updated with page param (?page=N). IntersectionObserver with root=listRef.current loads next page at bottom. Badge uses meta.total. Append vs replace logic on currentPage. RTK Query caching pitfall avoided.",
             "", "Week 5", "Mar 2026", "ROW"))

# ═══ SECTION 6 ─ Search & Filter ══════════════
rows.append(("", "🔍  6. Search & Filter Responsiveness", "", "", "", "", "", "", "", "", "", "SECTION"))
rows.append(("6.1", "Dashboard product search: input debouncing (500ms)",
             "Search", "✅ Done", "🔴 Critical", "S",
             "components/features/dashboard/Dashboard.tsx",
             "setTimeout 500ms debounce on search input; skips API call when empty",
             "", "Week 5", "Feb 2026", "ROW"))
rows.append(("6.2", "History search: input debouncing",
             "Search", "✅ Done", "🟠 High", "S",
             "components/features/history/History.tsx",
             "debouncedSearch pattern applied",
             "", "Week 5", "Feb 2026", "ROW"))
rows.append(("6.3", "Profitability calculator: debounced calculation on input change",
             "Search", "✅ Done", "🟠 High", "S",
             "components/features/dashboard/product/profitability-calculator.tsx",
             "lodash debounce 500ms applied to handleCalculateProfitability",
             "", "Week 5", "Feb 2026", "ROW"))
rows.append(("6.4", "Audit remaining search/filter inputs for missing debounce (UPC, Seller, GoCompare, Monitor)",
             "Search", "✅ Done", "🟠 High", "M",
             "components/features/upc-scanner/UpcScanner.tsx\ncomponents/features/seller/Seller.tsx\ncomponents/features/go-compare/\ncomponents/features/monitor-list/Monitor.tsx",
             "Audit complete — no changes needed. UPC: lodash debounce 300ms (client-side filter). Seller: custom debounce 500ms → API. GoCompare: all searches are manual button triggers. Monitor: client-side array filter only, no API call.",
             "", "Week 5", "Mar 2026", "ROW"))
rows.append(("6.5", "Integrate fuse.js for client-side fuzzy search on large local datasets",
             "Search", "🚫 Ignored", "🟡 Medium", "L",
             "components/features/upc-scanner/UpcScanner.tsx",
             "Assessed and deliberately excluded. Search filters user's own scan list (typically 5-30 items) — dataset too small for fuzzy search to add value. includes() filter is instant at this size. Fuse.js adds ~24KB dependency for negligible gain.",
             "", "Week 6", "", "ROW"))
rows.append(("6.6", "Move heavy reverse-search filtering to server side",
             "Search", "🚫 Ignored", "🟡 Medium", "XL",
             "app/api/go-compare/reverse-search/route.ts\napp/api/go-compare/reverse-search/categories/route.ts",
             "Assessed and deliberately excluded. Reverse search results are sent to the user's email — they are never displayed in-app. ReverseSearchTable component with client-side sorting exists but is not rendered anywhere. No in-app results table to optimise.",
             "", "Week 6", "", "ROW"))

# ═══ SECTION 7 ─ Skeleton Screens ══════════════
rows.append(("", "💀  7. Loading States & Skeleton Screens", "", "", "", "", "", "", "", "", "", "SECTION"))
rows.append(("7.1", "MarketAnalysis component: skeleton screen",
             "Skeleton", "✅ Done", "🟠 High", "S",
             "components/features/dashboard/product/market-analysis.tsx",
             "MarketAnalysisSkeleton component implemented (line 356)",
             "", "Week 6", "Feb 2026", "ROW"))
rows.append(("7.2", "ProfitabilityCalculator component: skeleton screen",
             "Skeleton", "✅ Done", "🟠 High", "S",
             "components/features/dashboard/product/profitability-calculator.tsx",
             "Skeleton active paragraph shown during load state",
             "", "Week 6", "Feb 2026", "ROW"))
rows.append(("7.3", "Monitor List: skeleton screen",
             "Skeleton", "✅ Done", "🟡 Medium", "S",
             "components/features/monitor-list/Monitor.tsx",
             "Skeleton.Button + Skeleton paragraph rows shown on load",
             "", "Week 6", "Feb 2026", "ROW"))
rows.append(("7.4", "Settings page: skeleton screen",
             "Skeleton", "✅ Done", "🟡 Medium", "S",
             "components/features/settings/Settings.tsx",
             "Skeleton active paragraph on load",
             "", "Week 6", "Feb 2026", "ROW"))
rows.append(("7.5", "Dashboard search results: replace spinner with skeleton cards",
             "Skeleton", "✅ Done", "🟠 High", "M",
             "components/features/dashboard/Dashboard.tsx",
             "ProductCardSkeleton: 6 skeleton rows each with 64px image pulse + 3 text lines (title w-3/4, ASIN/UPC w-1/2, category w-2/5) matching exact real card layout. No phantom right-side element. Replaces CircularLoader during isLoading/isFetching. Header row also skeletonised.",
             "", "Week 6", "Mar 2026", "ROW"))
rows.append(("7.6", "Seller page product cards: skeleton during load",
             "Skeleton", "✅ Done", "🟠 High", "M",
             "components/features/seller/Seller.tsx",
             "3 skeleton rows matching exact 3-col grid (items-stretch). Col 1: 166×197px image + icon pair + 6 text lines + stats bar. Col 2: header badge + 5 offer rows. Col 3: chart block. Grid columns match real card (sm:2 lg:[3fr_2fr_2fr] xl:[4fr_2fr_2fr]). Shows during productLoading/loading.",
             "", "Week 6", "Mar 2026", "ROW"))
rows.append(("7.7", "ProductDetails sub-panels: skeleton screens for BuyboxAnalysis, SalesAnalytics, Keepa",
             "Skeleton", "🚫 Ignored", "🟠 High", "L",
             "components/features/dashboard/product/buybox-analysis.tsx\ncomponents/features/dashboard/product/sales-analytics.tsx\ncomponents/features/dashboard/product/keepa.tsx",
             "Assessed and deliberately excluded. ProductDetails already has a main page loader that covers initial load; sub-panel skeletons add complexity without meaningful UX uplift.",
             "", "Week 6", "", "ROW"))
rows.append(("7.8", "UPC Scanner progress indicator (status bar for long-running scans)",
             "Skeleton", "🚫 Ignored", "🟡 Medium", "M",
             "components/features/upc-scanner/UpcScanner.tsx",
             "Assessed and deliberately excluded. products_found represents matched products, not items processed — not a reliable progress fraction. API exposes no processed_count or explicit progress field. Manual refresh button already covers the use case.",
             "", "Week 7", "", "ROW"))

# ═══ SECTION 8 ─ Offline Support ══════════════
rows.append(("", "📴  8. Offline Support & Resilience", "", "", "", "", "", "", "", "", "", "SECTION"))
rows.append(("8.1", "Cache recent product data locally (last 10-20 ASINs viewed)",
             "Offline", "🚫 Ignored", "🟡 Medium", "XL",
             "utils/ (new: localProductCache.ts)",
             "Assessed and deliberately excluded. RTK Query in-session cache already handles same-session navigation. Product data (prices, buybox, rankings) changes hourly — cross-session caching risks showing stale sourcing data. Backend search history API already covers re-finding viewed products. XL effort for negligible gain.",
             "", "Week 7", "", "ROW"))
rows.append(("8.2", "Cache user preferences locally (marketplace, date range, calculator inputs)",
             "Offline", "✅ Done", "🟡 Medium", "M",
             "redux/slice/globalSlice.ts\ncomponents/ui/CountrySelect.tsx",
             "globalSlice initialState reads localStorage key 'optisage-marketplace' on mount (typeof window guard for SSR). CountrySelect.handleSelect writes {marketplaceId, currencyCode, currencySymbol} to localStorage on selection. First-time users get hardcoded defaults (Canada/CAD/C$); returning users restore their last marketplace selection.",
             "", "Week 7", "Mar 2026", "ROW"))
rows.append(("8.3", "Edge case: show offline banner + degraded UI when network is lost",
             "Offline", "✅ Done", "🟢 Low", "M",
             "app/(dashboard)/layout.tsx",
             "isOnline state seeded from navigator.onLine on mount. window online/offline event listeners update state. Amber banner renders below DashNav when offline: warning icon + stale data message. Auto-dismisses on reconnect. RTK Query refetchOnReconnect:true handles data refresh.",
             "", "Week 8", "Mar 2026", "ROW"))
rows.append(("8.4", "Service Worker / PWA manifest (optional advanced offline)",
             "Offline", "🚫 Ignored", "🟢 Low", "XL",
             "public/manifest.json",
             "Assessed and deliberately excluded. Out of scope for current optimisation phase. Can be revisited if PWA/offline-first is added to product roadmap.",
             "", "Week 8+", "", "ROW"))
rows.append(("8.5", "Custom 404 page (not-found.tsx) with branded design and navigation shortcuts",
             "Offline", "✅ Done", "🟠 High", "S",
             "app/not-found.tsx",
             "Server component. Logo + #18CB96 badge + friendly headline + 3 nav shortcuts (Dashboard, Go Compare, Home). Matches Tailwind design system. Replaces Next.js default 404.",
             "", "Week 5", "Mar 2026", "ROW"))
rows.append(("8.6", "Custom network / global error page (error.tsx + global-error.tsx)",
             "Offline", "✅ Done", "🟠 High", "S",
             "app/error.tsx\napp/global-error.tsx",
             "error.tsx: client component, Sentry capture, Try Again + Dashboard + Home links, branded with logo. global-error.tsx: inline styles (no Tailwind/Next.js Image in root layout replacement), inline SVG wordmark, retry-only CTA, Sentry capture.",
             "", "Week 5", "Mar 2026", "ROW"))

# ─────────────────────────────────────────────
# WRITE ROWS
# ─────────────────────────────────────────────
current_row = 3
for i, row_data in enumerate(rows):
    is_phase  = row_data[-1] == "PHASE"
    is_sec    = row_data[-1] == "SECTION"
    is_data   = row_data[-1] == "ROW"
    alt       = (i % 2 == 0)
    write_row(ws, current_row, list(row_data[:11]), is_phase_header=is_phase, is_section=is_sec, alt=alt)
    if is_phase:
        ws.merge_cells(f"B{current_row}:K{current_row}")
        ws.row_dimensions[current_row].height = 26
    elif is_sec:
        ws.merge_cells(f"B{current_row}:K{current_row}")
        ws.row_dimensions[current_row].height = 20
    current_row += 1

# ─────────────────────────────────────────────
# SHEET 2 — SUMMARY DASHBOARD
# ─────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")
ws2.column_dimensions["A"].width = 30
ws2.column_dimensions["B"].width = 14
ws2.column_dimensions["C"].width = 14
ws2.column_dimensions["D"].width = 14
ws2.column_dimensions["E"].width = 14
ws2.column_dimensions["F"].width = 22

# Title
ws2.merge_cells("A1:F1")
t = ws2["A1"]
t.value = "Horizon FE Optimization — Progress Summary"
t.font = Font(name="Calibri", bold=True, size=14, color=C_WHITE)
t.fill = fill(C_HEADER_DARK)
t.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 32

# Header row
summary_headers = ["Area", "✅ Done", "⚠️ Partial", "❌ Not Done", "📋 Planned", "% Complete"]
for ci, sh in enumerate(summary_headers, 1):
    c = ws2.cell(row=2, column=ci, value=sh)
    c.font = Font(name="Calibri", bold=True, size=10, color=C_WHITE)
    c.fill = fill(C_HEADER_MID)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = thin_border()
ws2.row_dimensions[2].height = 20

# Count statuses per area from rows data
from collections import defaultdict
area_counts = defaultdict(lambda: {"✅ Done": 0, "⚠️ Partial": 0, "❌ Not Done": 0, "📋 Planned": 0, "🚫 Ignored": 0})
for row_data in rows:
    if row_data[-1] == "ROW":
        area  = str(row_data[2])
        status = str(row_data[3])
        if status in area_counts[area]:
            area_counts[area][status] += 1

summary_rows = [
    ("Folder Restructure",  area_counts["Architecture"]["✅ Done"],  area_counts["Architecture"]["⚠️ Partial"],  area_counts["Architecture"]["❌ Not Done"],  area_counts["Architecture"]["📋 Planned"]),
    ("Virtual Scrolling",   area_counts["Virtual Scroll"]["✅ Done"], area_counts["Virtual Scroll"]["⚠️ Partial"], area_counts["Virtual Scroll"]["❌ Not Done"], area_counts["Virtual Scroll"]["📋 Planned"]),
    ("API Caching",         area_counts["API Caching"]["✅ Done"],    area_counts["API Caching"]["⚠️ Partial"],    area_counts["API Caching"]["❌ Not Done"],    area_counts["API Caching"]["📋 Planned"]),
    ("Chart Rendering",     area_counts["Charts"]["✅ Done"],         area_counts["Charts"]["⚠️ Partial"],         area_counts["Charts"]["❌ Not Done"],         area_counts["Charts"]["📋 Planned"]),
    ("Code Splitting",      area_counts["Code Splitting"]["✅ Done"], area_counts["Code Splitting"]["⚠️ Partial"], area_counts["Code Splitting"]["❌ Not Done"], area_counts["Code Splitting"]["📋 Planned"]),
    ("Loading UX",          area_counts["Loading UX"]["✅ Done"],     area_counts["Loading UX"]["⚠️ Partial"],     area_counts["Loading UX"]["❌ Not Done"],     area_counts["Loading UX"]["📋 Planned"]),
    ("Search / Filter",     area_counts["Search"]["✅ Done"],         area_counts["Search"]["⚠️ Partial"],         area_counts["Search"]["❌ Not Done"],         area_counts["Search"]["📋 Planned"]),
    ("Skeleton Screens",    area_counts["Skeleton"]["✅ Done"],       area_counts["Skeleton"]["⚠️ Partial"],       area_counts["Skeleton"]["❌ Not Done"],       area_counts["Skeleton"]["📋 Planned"]),
    ("Offline / Resilience",area_counts["Offline"]["✅ Done"],        area_counts["Offline"]["⚠️ Partial"],        area_counts["Offline"]["❌ Not Done"],        area_counts["Offline"]["📋 Planned"]),
]

for ri, (area, done, partial, notdone, planned) in enumerate(summary_rows, 3):
    total = done + partial + notdone + planned
    pct = round(((done + 0.5 * partial) / total * 100) if total > 0 else 0, 1)

    bg = C_ROW_ALT if ri % 2 == 0 else C_WHITE
    vals = [area, done, partial, notdone, planned, f"{pct}%"]
    for ci, v in enumerate(vals, 1):
        c = ws2.cell(row=ri, column=ci, value=v)
        c.font = Font(name="Calibri", size=10)
        c.fill = fill(bg)
        c.border = thin_border()
        c.alignment = Alignment(horizontal="center" if ci > 1 else "left", vertical="center")
        if ci == 2 and v:
            c.fill = fill(C_DONE); c.font = Font(name="Calibri", bold=True, size=10, color=C_DONE_FONT)
        if ci == 3 and v:
            c.fill = fill(C_PARTIAL); c.font = Font(name="Calibri", bold=True, size=10, color=C_PARTIAL_FONT)
        if ci == 4 and v:
            c.fill = fill(C_TODO); c.font = Font(name="Calibri", bold=True, size=10, color=C_TODO_FONT)
        if ci == 5 and v:
            c.fill = fill(C_PLANNED); c.font = Font(name="Calibri", bold=True, size=10, color=C_PLANNED_FONT)
    ws2.row_dimensions[ri].height = 18

# Totals row
total_row = len(summary_rows) + 3
all_done     = sum(r[1] for r in summary_rows)
all_partial  = sum(r[2] for r in summary_rows)
all_notdone  = sum(r[3] for r in summary_rows)
all_planned  = sum(r[4] for r in summary_rows)
all_total    = all_done + all_partial + all_notdone + all_planned
all_pct      = round(((all_done + 0.5 * all_partial) / all_total * 100) if all_total > 0 else 0, 1)

totals_vals = ["TOTAL", all_done, all_partial, all_notdone, all_planned, f"{all_pct}%"]
for ci, v in enumerate(totals_vals, 1):
    c = ws2.cell(row=total_row, column=ci, value=v)
    c.font = Font(name="Calibri", bold=True, size=11, color=C_WHITE)
    c.fill = fill(C_HEADER_DARK)
    c.alignment = Alignment(horizontal="center" if ci > 1 else "left", vertical="center")
    c.border = thin_border()
ws2.row_dimensions[total_row].height = 22

# Legend section
legend_row = total_row + 2
ws2.cell(row=legend_row, column=1, value="LEGEND").font = Font(name="Calibri", bold=True, size=11)
legend_items = [
    ("✅ Done",      C_DONE,    C_DONE_FONT,    "Fully implemented and verified"),
    ("⚠️ Partial",  C_PARTIAL, C_PARTIAL_FONT, "Started or partially implemented — needs completion"),
    ("❌ Not Done",  C_TODO,    C_TODO_FONT,    "Not yet started"),
    ("📋 Planned",  C_PLANNED, C_PLANNED_FONT, "Deferred / conditional — in backlog"),
    ("🚫 Ignored",  "E2E8F0",  "64748B",       "Assessed and deliberately excluded — not worth the complexity or effort"),
]
for i, (label, bg_hex, fg_hex, desc) in enumerate(legend_items):
    r = legend_row + 1 + i
    c1 = ws2.cell(row=r, column=1, value=label)
    c1.fill = fill(bg_hex)
    c1.font = Font(name="Calibri", bold=True, size=10, color=fg_hex)
    c1.border = thin_border()
    c1.alignment = Alignment(horizontal="center", vertical="center")
    c2 = ws2.cell(row=r, column=2, value=desc)
    ws2.merge_cells(f"B{r}:F{r}")
    c2.font = Font(name="Calibri", size=10)
    c2.alignment = Alignment(vertical="center")
    c2.border = thin_border()

# Effort legend
eff_row = legend_row + len(legend_items) + 2
ws2.cell(row=eff_row, column=1, value="EFFORT KEY").font = Font(name="Calibri", bold=True, size=11)
effort_items = [("XS", "< 1 hour"), ("S", "Half day"), ("M", "1-2 days"), ("L", "3-5 days"), ("XL", "> 1 week")]
for i, (ef, desc) in enumerate(effort_items):
    r = eff_row + 1 + i
    c1 = ws2.cell(row=r, column=1, value=ef)
    c1.font = Font(name="Calibri", bold=True, size=10)
    c1.border = thin_border()
    c1.alignment = Alignment(horizontal="center", vertical="center")
    c2 = ws2.cell(row=r, column=2, value=desc)
    ws2.merge_cells(f"B{r}:F{r}")
    c2.font = Font(name="Calibri", size=10)
    c2.alignment = Alignment(vertical="center")
    c2.border = thin_border()

# ─────────────────────────────────────────────
# SHEET 3 — NEXT ACTIONS
# ─────────────────────────────────────────────
ws3 = wb.create_sheet("Next Actions")
ws3.column_dimensions["A"].width = 6
ws3.column_dimensions["B"].width = 42
ws3.column_dimensions["C"].width = 18
ws3.column_dimensions["D"].width = 14
ws3.column_dimensions["E"].width = 14
ws3.column_dimensions["F"].width = 36

ws3.merge_cells("A1:F1")
t3 = ws3["A1"]
t3.value = "Next Actions — Immediate Implementation Queue"
t3.font = Font(name="Calibri", bold=True, size=14, color=C_WHITE)
t3.fill = fill(C_HEADER_DARK)
t3.alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[1].height = 32

na_headers = ["#", "Task", "Area", "Effort", "Priority", "Acceptance Criteria"]
for ci, h in enumerate(na_headers, 1):
    c = ws3.cell(row=2, column=ci, value=h)
    c.font = Font(name="Calibri", bold=True, size=10, color=C_WHITE)
    c.fill = fill(C_HEADER_MID)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = thin_border()
ws3.row_dimensions[2].height = 20

next_actions = [
    ("—", "All planned optimisation items have been completed or assessed.",
     "—", "—", "—",
     "The tracker is fully resolved. Review for any newly discovered items or add new work streams as needed."),
]

for ri, (num, task, area, effort, priority, ac) in enumerate(next_actions, 3):
    alt = ri % 2 == 0
    bg = C_ROW_ALT if alt else C_WHITE
    vals = [num, task, area, effort, priority, ac]
    for ci, v in enumerate(vals, 1):
        c = ws3.cell(row=ri, column=ci, value=v)
        c.fill = fill(bg)
        c.border = thin_border()
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.font = Font(name="Calibri", size=9)
        if ci == 1:
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.font = Font(name="Calibri", bold=True, size=10)
        if ci == 5 and v in {"🔴 Critical", "🟠 High"}:
            bg_p = "FECACA" if "Critical" in v else "FED7AA"
            fg_p = "991B1B" if "Critical" in v else "9A3412"
            c.fill = fill(bg_p)
            c.font = Font(name="Calibri", bold=True, size=9, color=fg_p)
            c.alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[ri].height = 44

# ─────────────────────────────────────────────
# SAVE
# ─────────────────────────────────────────────
output_path = "docs/horizon-fe-optimization-tracker.xlsx"
wb.save(output_path)
print(f"Tracker saved -> {output_path}")
